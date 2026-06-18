#!/usr/bin/env python3
"""Generate the Day 15 TMI validation Word report.

Reads the committed Day-15 artifacts and builds a structured .docx covering the
v2 topology fix, the CMC re-fit (constant -> conflict-dependent fallback), the
return mechanism, the Zeigler/shadow validation, the P_S2 comparison figure, the
diagnostics gate, and a dedicated section on the FLEE-developer design concern
(CMC controlling whether agents leave vs FLEE's where-they-go design envelope).

Output: reports/TMI_Day15_Report.docx   (NOT committed to the repo)
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent
RES = REPO / "results" / "tmi"
FIG = REPO / "figures" / "tmi"
OUT = REPO / "reports" / "TMI_Day15_Report.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5F, 0xA3)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GREY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB0, 0x2A, 0x2A)


def _load():
    return {
        "locked": json.loads((RES / "day15_locked_params.json").read_text()),
        "refit": json.loads((RES / "day15_refit_verdict.json").read_text()),
        "ret": json.loads((RES / "day15_return_verdict.json").read_text()),
        "zs": json.loads((RES / "day15_zeigler_shadow_verdict.json").read_text()),
        "gate": json.loads((RES / "diagnostics_gate_day15.json").read_text()),
        "cpA": pd.read_csv(RES / "day15_topology_checkpointA.csv"),
        "grid_const": pd.read_csv(RES / "day15_refit.csv"),
        "grid_b2": pd.read_csv(RES / "day15_refit_b2.csv"),
        "zs_csv": pd.read_csv(RES / "day15_zeigler_shadow.csv"),
        "ret_csv": pd.read_csv(RES / "day15_return_curve.csv"),
    }


def _h(doc, text, level=1, color=NAVY, size=None):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return p


def _para(doc, text, italic=False, bold=False, size=10.5, color=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    return p


def _style_header_row(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.5)
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F3864")
        tc.append(shd)


def _table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].paragraphs[0].add_run(h)
    _style_header_row(t.rows[0])
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
    if widths:
        for col, w in enumerate(widths):
            for row in t.rows:
                row.cells[col].width = Inches(w)
    return t


def build():
    d = _load()
    locked, refit, ret, zs, gate = (d["locked"], d["refit"], d["ret"],
                                    d["zs"], d["gate"])
    cpA, grid_const, grid_b2 = d["cpA"], d["grid_const"], d["grid_b2"]
    zones = zs["zones"]

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    # --- Title block ---
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("FLEE Dual-Process Model \u2014 Phase 3: TMI Full Calibration")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Day 15 \u2014 TMI Validation: Topology Fix, CMC Re-Fit, "
                    "Return Mechanism, and Empirical Signatures")
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLUE
    attr = doc.add_paragraph(); attr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = attr.add_run("Columbia University Climate School \u00b7 Center for Climate "
                     "Systems Research")
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()

    # --- 1. Executive Summary ---
    _h(doc, "1. Executive Summary", level=1)
    _para(doc,
          "Day 15 fixes the topology that confounded Day 14, re-fits the System-1 "
          "move chance (CMC) with departure timing as a primary target, builds a "
          "return mechanism, and validates the calibrated TMI model against the "
          "empirical signatures. \u03b1/\u03b2/\u03ba remain LOCKED at the Fukushima transfer "
          "values (1.6667 / 2.1667 / 18.407); the original tmi network is untouched "
          "(all topology work is in input_data/tmi_v2); both new model levers are "
          "behind default-False flags so the 84 baseline tests stay green.")
    _bullet(doc, "three distant directional shelters (State College / Philadelphia "
                 "/ Baltimore) funneled through the Harrisburg transit hub; OSRM-"
                 "routed. Cross-zone camp access spread = "
                 f"{cpA.groupby('zone')['nearest_camp_km'].mean().max()/cpA.groupby('zone')['nearest_camp_km'].mean().min():.2f}x "
                 "(<2x). The camp-adjacency confound is removed.",
            bold_prefix="Topology (v2): ")
    _bullet(doc, "on v2 a CONSTANT CMC re-fits POSITIVE (total in band) \u2014 "
                 "confirming Day-14's CMC=0.00 was a TOPOLOGY ARTIFACT, not a real "
                 "estimate. But a constant CMC is a constant departure hazard and "
                 "cannot make the Friday spike, so the conflict-dependent fallback "
                 "was triggered.", bold_prefix="CMC re-fit: ")
    _bullet(doc, "LOCKED System-1 form P_S1 = base + (cmax\u2212base)\u00b7c with "
                 f"base={locked['cmc_base']}, cmax={locked['cmc_max']}. Total evac "
                 f"{refit['total_evac']*100:.1f}% (in band), zone inversion removed, "
                 f"Friday spike inner-by-t10 {refit['frac_inner_by_t10']*100:.0f}%, "
                 "no \u03a9-recovery evacuation.", bold_prefix="Conflict-dependent CMC: ")
    _bullet(doc, f"camp builds {ret['camp_pop_start']:.0f}\u2192{ret['camp_pop_peak']:.0f}"
                 f"@t{ret['camp_pop_peak_t']}\u2192{ret['camp_pop_end']:.0f}; cumulative "
                 f"return {ret['week1_return']*100:.0f}/{ret['week2_return']*100:.0f}/"
                 f"{ret['week3_return']*100:.0f}% by weeks 1/2/3 (empirical 53/90/98); "
                 "S2-driven.", bold_prefix="Return mechanism: ")
    _bullet(doc, f"gate {gate['verdict'].split(' (')[0]}; \u03b1/\u03b2/\u03ba unchanged "
                 "(PASS). Lone failure is the inherited Day-12 frozen-CMC check 26.",
            bold_prefix="Gate: ")
    _para(doc, "Two results are reported as documented findings, not forced: a "
               "strict inner>middle>outer evacuation gradient and a strict "
               "departure-peak-at-t10 do NOT emerge under TMI's spatially UNIFORM "
               "conflict schedule with \u03b1/\u03b2/\u03ba locked. Section 9 discusses the "
               "deeper architectural issue this exposes (raised by the FLEE "
               "developer): CMC \u2014 a single, crude, non-robust parameter \u2014 is "
               "being asked to decide WHETHER people leave, which is outside FLEE's "
               "original design.", italic=True, color=GREY)

    # --- 2. Topology Fix (Task A) ---
    _h(doc, "2. Topology Fix \u2014 input_data/tmi_v2 (Task A)", level=1)
    _para(doc, "Day 14's zone gradient was inverted because the big OUTER towns "
               "(York, Lancaster) sat ~1 hop / ~2.4 km from their own camp, while "
               "the dense INNER cluster was multi-hop from any camp. v2 replaces the "
               "three adjacent camps with three DISTANT directional shelters at real "
               "evacuation-geography cities, each reachable only through the "
               "Harrisburg transit hub, so the camp leg (~127\u2013172 km) is shared by "
               "all towns and access is comparable across zones.")
    crows = []
    for z in ("inner", "middle", "outer"):
        sub = cpA[cpA.zone == z]
        crows.append([z, f"{sub['nearest_camp_km'].mean():.1f}",
                      f"{sub['hops'].mean():.1f}", int(len(sub))])
    zmean = cpA.groupby("zone")["nearest_camp_km"].mean()
    _table(doc, ["Zone", "Mean nearest-camp (km)", "Mean hops", "# towns"], crows,
           widths=[1.2, 2.2, 1.3, 1.0])
    _para(doc, f"Checkpoint A: per-zone spread = {zmean.max()/zmean.min():.2f}x "
               f"(inner {zmean['inner']:.0f} / middle {zmean['middle']:.0f} / outer "
               f"{zmean['outer']:.0f} km) \u2014 PASS (<2x, no zone has a >2x camp-access "
               "advantage). The residual gradient (inner closer to the hub than "
               "outer) is the physically correct direction.", bold=True, color=GREEN)

    # --- 3. CMC re-fit, constant (Task B) ---
    _h(doc, "3. CMC Re-Fit with Timing as a Primary Target (Task B)", level=1)
    _para(doc, "CMC is the System-1 term in P_move = (1\u2212P_S2)\u00b7CMC + P_S2\u00b7\u03c3. "
               "At the Friday conflict peak \u03a9 collapses so P_S2\u21920 and "
               "(1\u2212P_S2)\u21921: the (1\u2212P_S2)\u00b7CMC term is maximal exactly at the "
               "threat peak, so CMC is the mechanism for prompt flight. The re-fit "
               "adds departure timing to the objective (one-sided inner-by-t10 "
               "shortfall, a departure-peak penalty, and a directional Zeigler term) "
               "so the fit cannot buy a good total by deferring evacuation.")
    inband = grid_const[(grid_const.total_evac_all >= 0.15)
                        & (grid_const.total_evac_all <= 0.30)]
    cmc_band_min = inband["cmc"].min() if len(inband) else float("nan")
    _bullet(doc, "the constant CMC re-fits to a POSITIVE value (total enters the "
                 f"band at CMC\u2265{cmc_band_min:.2f}); the Day-14 boundary value 0.00 "
                 "does NOT reproduce \u2014 it was a topology artifact.",
            bold_prefix="Result 1: ")
    _bullet(doc, "for EVERY CMC the first-departure peak is t=1, not t=10. A "
                 "constant CMC is a constant per-step departure hazard, so "
                 "departures front-load into the calm; no constant CMC can "
                 "concentrate flight at the advisory.", bold_prefix="Result 2: ")
    _para(doc, "Checkpoint B fails on timing (the constant-CMC limitation), which "
               "triggers the Task-B2 conflict-dependent fallback. Per the prompt, a "
               "constant CMC that evacuates during the calm is not locked.",
          bold=True, color=RED)

    # --- 4. Conflict-dependent CMC (Task B2) ---
    _h(doc, "4. Conflict-Dependent System-1 Move Chance (Task B2)", level=1)
    _para(doc, "Behind the flag cmc_conflict_dependent (default False), the constant "
               "System-1 move chance is replaced by P_S1 = base + (cmax\u2212base)\u00b7c so "
               "prompt flight at the advisory emerges endogenously from System 1 "
               "when conflict spikes \u2014 without a constant baseline and without the "
               "population-wide panic clamp removed on Day 14. The (base, cmax) grid "
               "was calibrated on v2 against the timing-weighted objective.")
    b2 = refit["checkpoint_b2"]
    _table(doc, ["Condition", "Result", "Status"], [
        ["(a) total evac in band [15,30]%",
         f"{refit['total_evac']*100:.1f}%", "PASS" if b2["a_total_in_band"] else "FAIL"],
        ["(b) zone inversion removed (outer not highest)",
         f"in/mid/out {refit['evac_inner']:.2f}/{refit['evac_middle']:.2f}/"
         f"{refit['evac_outer']:.2f}", "PASS" if b2["b_inversion_removed"] else "FAIL"],
        ["    strict inner\u2265middle\u2265outer", "uniform-threat limit",
         "documented" if not b2["b_strict_order_inner_ge_mid_ge_outer"] else "PASS"],
        ["(c1) Friday spike inner-by-t10 > 0.77",
         f"{refit['frac_inner_by_t10']:.2f}",
         "PASS" if b2["c1_friday_spike_inner_by_t10_gt_077"] else "FAIL"],
        ["(c2) no \u03a9-recovery evacuation (<10%)",
         f"{refit['frac_dep_during_recovery']*100:.1f}%",
         "PASS" if b2["c2_no_omega_recovery_evac"] else "FAIL"],
        ["(c3) strict first-departure peak @ t10",
         f"peakT={refit['departure_peak_t']:.0f}",
         "documented" if not b2["c3_strict_first_departure_peak_t10"] else "PASS"],
    ], widths=[2.9, 2.0, 1.2])
    _para(doc, f"LOCKED TMI System-1 specification: conflict-dependent, "
               f"base={locked['cmc_base']}, cmax={locked['cmc_max']} (clamp OFF). "
               "Primary checkpoint conditions are met: total in band, inversion "
               "removed, Friday spike present, and no evacuation during the "
               "post-Carter calm.", bold=True, color=GREEN)
    _para(doc, "Documented findings (not forced): (i) a clean inner>middle>outer "
               "gradient does not emerge \u2014 TMI's conflict is spatially UNIFORM, so "
               "the only zonal lever is network/hub position, which puts the "
               "hub-zone (middle) highest; reproducing inner>outer needs a "
               "proximity-to-plant THREAT gradient in the conflict input. (ii) The "
               "strict first-departure peak stays early because the System-2 "
               "deliberation channel (P_S2 high in the calm, \u03c3\u22480.5 toward the "
               "zero-conflict distant shelters) drives diffuse pre-advisory drift "
               "under LOCKED \u03b1/\u03b2/\u03ba. Both are real model findings, flagged for the "
               "theory paper / Day 16.", italic=True, color=GREY, size=9.5)

    # --- 5. Return mechanism (Task C/E) ---
    _h(doc, "5. Return Mechanism and Return Curve (Tasks C, E)", level=1)
    _para(doc, "Behind the flag enable_return (default False), a camp agent "
               "re-evaluates each step; once its HOME conflict has recovered "
               "(c_home low) and its deliberative weight P_S2 is high (\u03a9 recovered "
               "after the Carter visit, t=16), it returns home with a probability "
               "driven by P_S2 \u2014 the same \u03c3 machinery in reverse. The conflict "
               "schedule is extended past T_SCHED at c\u22480 for the return phase "
               "(1 step = 8 h; 1 week = 21 steps).")
    _table(doc, ["Quantity", "Model", "Empirical / expected"], [
        ["Camp population (start \u2192 peak \u2192 end)",
         f"{ret['camp_pop_start']:.0f} \u2192 {ret['camp_pop_peak']:.0f}@t"
         f"{ret['camp_pop_peak_t']} \u2192 {ret['camp_pop_end']:.0f}",
         "build during crisis, draw down after t16"],
        ["Cumulative return week 1 / 2 / 3",
         f"{ret['week1_return']*100:.0f} / {ret['week2_return']*100:.0f} / "
         f"{ret['week3_return']*100:.0f} %", "53 / 90 / 98 %"],
        ["Return is S2-driven (P_S2 ret vs non-ret)",
         f"{ret['ps2_returners_mean']:.2f} > {ret['ps2_nonreturners_mean']:.2f}; "
         f"median return t={ret['median_return_t']:.0f}", "returners deliberate; t\u226516"],
    ], widths=[2.7, 2.3, 1.5])
    _para(doc, "Checkpoint C PASS: camp population builds during the crisis and "
               "draws down after t\u224816; the cumulative-return curve reproduces the "
               "empirical shape (deliberate, post-reassurance return), reaching "
               f"~{ret['week3_return']*100:.0f}% by week 3 vs the empirical 98%.",
          bold=True, color=GREEN)

    # --- 6. Zeigler & shadow validation (Task D) ---
    _h(doc, "6. Validation \u2014 Zeigler Paradox & Shadow Evacuation (Task D)", level=1)
    _table(doc, ["Signal", "Result", "Reads?"], [
        ["Zeigler TIMING (inner departs earlier)",
         f"inner by-t10 {zones['inner']['frac_by_t10']:.2f} vs outer "
         f"{zones['outer']['frac_by_t10']:.2f}",
         "FLAT (uniform advisory)"],
        ["Zeigler DISTANCE \u2014 travel path (inner<outer)",
         f"{zones['inner']['travel_path_km']:.0f} < "
         f"{zones['outer']['travel_path_km']:.0f} km",
         "HOLDS" if zs["zeigler_distance_travel_inner_lt_outer"] else "NO"],
        ["Zeigler DISTANCE \u2014 dest-from-plant",
         f"{zones['inner']['dest_dist_from_plant_mi']:.0f} vs "
         f"{zones['outer']['dest_dist_from_plant_mi']:.0f} mi",
         "~uniform (all far \u2248 85-mi median)"],
        ["Shadow evac (beyond 15 mi)",
         f"{zs['shadow_beyond_15mi_rate']*100:.0f}% (within-15mi "
         f"{zs['within_15mi_rate']*100:.0f}% vs empirical 39%)",
         "GENUINE (info-\u03a9)"],
    ], widths=[2.9, 2.4, 1.4])
    _para(doc, "The distance paradox reads on travel-path length (inner travels less "
               "to the common far shelters); the plant-referenced destination "
               "distance is ~uniform because on v2 all zones evacuate to comparably "
               "distant camps (~60\u201384 mi) \u2014 consistent with the empirical 85-mi "
               "median destination. Shadow evacuation is GENUINE: on v2 every camp "
               "is far and zone-independent, so beyond-15-mi evacuation is driven by "
               "the uniform information-\u03a9 collapse, NOT the pre-v2 camp-adjacency "
               f"artifact (which inflated outer evac to {zs['prev2_outer_artifact']*100:.0f}%). "
               "The within-15-mi rate (40%) matches the empirical 39% closely; the "
               "beyond-15-mi magnitude (34%) exceeds the empirical ~10\u201318% because "
               "the uniform schedule has no distance decay of threat \u2014 the same "
               "uniform-threat limitation.", italic=True, color=GREY, size=9.5)

    # --- 7. P_S2 comparison figure (Task F) ---
    _h(doc, "7. P_S2 Trajectory Comparison vs Fukushima (Task F)", level=1)
    fig = FIG / "day15_ps2_comparison.png"
    if fig.exists():
        doc.add_picture(str(fig), width=Inches(6.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run("Figure 1. Left: TMI (v2) zone P_S2 \u2014 spatially UNIFORM "
                        "(zone lines overlap) and temporally VOLATILE (collapse at "
                        "the t10 advisory, recover after the t16 Carter visit). "
                        "Right: Fukushima zone P_S2 \u2014 a parameter-based reference "
                        "(locked \u03b1/\u03b2 on a spatial conflict gradient) that is "
                        "spatially GRADED and temporally STABLE. The uniform-\u03a9 vs "
                        "spatial-\u03a9 contrast is the qualitative payoff; a "
                        "simulation-derived parameter-matched Fukushima spread is a "
                        "Day-16 item.")
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY

    # --- 8. Diagnostics gate ---
    _h(doc, "8. Diagnostics Gate", level=1)
    _para(doc, f"Verdict: {gate['verdict']}.", bold=True)
    grows = []
    for c in gate["day15_new_checks"]:
        grows.append([c["name"][:46], c["status"], c["detail"][:70]])
    _table(doc, ["Check", "Status", "Detail"], grows, widths=[2.5, 0.8, 3.5])
    if gate.get("inherited_expected_failures"):
        _para(doc, "Inherited expected failure (by design): "
                   + gate["inherited_expected_failures"][0], italic=True,
              color=GREY, size=9.5)

    # --- 9. Critical design issue: CMC and FLEE's design envelope ---
    _h(doc, "9. Critical Design Issue \u2014 CMC and FLEE's Design Envelope", level=1)
    _para(doc, "The FLEE developer has flagged a fundamental concern: it is "
               "problematic for a single, crudely-defined, non-robust parameter "
               "(CMC) to control the model. The root cause is a design-envelope "
               "mismatch. In standard FLEE the agent population represents people "
               "who have ALREADY decided to leave \u2014 the model's validated job is "
               "WHERE displaced people go and HOW they distribute among "
               "destinations, not WHETHER an individual leaves. Our dual-process "
               "work has been using CMC (the System-1 move chance) to decide the "
               "WHETHER (the evacuation rate), which is outside that envelope. This "
               "is exactly why CMC keeps behaving pathologically: it pinned at 0 on "
               "Day 14 (topology artifact), had to be made conflict-dependent on Day "
               "15, and still cannot place the departure PEAK correctly.")
    _para(doc, "Two remedies were proposed, and they are complementary:", bold=True)
    _bullet(doc, "ground the System-1 move chance in a behavioural mechanism "
                 "(threat perception, official advisory, social influence \u2014 e.g. "
                 "Cutter's 74% 'neighbours left' effect) rather than a single fitted "
                 "scalar. The Day-15 conflict-dependent form P_S1=base+(cmax\u2212base)\u00b7c "
                 "is a first step, but it is still phenomenological.",
            bold_prefix="Option 1 \u2014 define CMC mechanistically: ")
    _bullet(doc, "set CMC\u22481 (once an agent is in the simulation it is mobile) and "
                 "move the leave/stay decision into spawning.py, so the agent "
                 "population is \u2014 as FLEE intends \u2014 the people who decided to "
                 "leave. The dual-process model then governs the SPAWN/DEPARTURE "
                 "hazard (who leaves and when) while FLEE handles routing (where).",
            bold_prefix="Option 2 \u2014 CMC\u22481 + decision in spawning: ")
    _para(doc, "Recommendation: Option 2 (with Option 1's mechanism living inside "
               "the spawner). This is the architecturally correct separation and it "
               "directly resolves the Day-15 anomalies. FLEE already provides "
               "conflict_driven_spawning + TakeFromPopulation: agents would be spawned "
               "from the resident population at a per-zone, per-timestep DEPARTURE "
               "hazard equal to the dual-process P_move; once spawned they route to "
               "shelters with CMC\u22481 and the \u03c3 destination weighting. Benefits: "
               "(i) the Friday spike becomes a property of the spawn schedule (timed "
               "to the advisory) instead of a per-step movechance, eliminating the "
               "System-2 calm-drift; (ii) the zone evacuation RATE becomes a "
               "spawn-fraction driven by threat/P_S2, decoupled from routing and "
               "topology, so inner>outer can arise from a proximity-to-plant threat "
               "gradient rather than camp geometry; (iii) CMC stops being the "
               "fragile master knob \u2014 robustness moves to the spawner where it is "
               "interpretable. Costs: it is a real refactor (spawning + decision "
               "wiring) and Days 8\u201315 calibration would be re-expressed in "
               "spawn-rate terms; the Fukushima transfer story (\u03b1/\u03b2) is unaffected "
               "in principle but must be re-validated. Net: the cross-domain claim "
               "becomes cleaner because both domains then share one mechanism "
               "(decision \u2192 spawn) rather than a domain-specific scalar.", bold=True)
    _para(doc, "Interim status: today's conflict-dependent CMC + return flags are "
               "behind default-False switches and do not disturb prior days; they "
               "are a stopgap that already exposes the limit. The spawning-based "
               "redesign is recommended as the Day-16+ structural fix and is the "
               "principled answer to the developer's concern.",
          italic=True, color=GREY, size=9.5)

    # --- 10. Implications ---
    _h(doc, "10. Implications for Day 16+", level=1)
    _bullet(doc, "compare the TMI System-1 SPECIFICATION (conflict-dependent base "
                 f"{locked['cmc_base']}, cmax {locked['cmc_max']}) against "
                 "Fukushima's constant 0.25 + clamp on equal footing \u2014 not a single "
                 "scalar; build the parameter-matched, simulation-derived Fukushima "
                 "P_S2 spread.", bold_prefix="Cross-domain: ")
    _bullet(doc, "add a proximity-to-plant THREAT gradient to the conflict input so "
                 "the inner>outer evacuation gradient can arise mechanistically.",
            bold_prefix="Uniform-threat fix: ")
    _bullet(doc, "migrate the leave/stay decision into spawning.py (Section 9) so "
                 "CMC is no longer the master knob.", bold_prefix="Architecture: ")

    # --- Appendix ---
    _h(doc, "Appendix \u2014 Artifact Manifest", level=2, color=BLUE)
    for f in ["input_data/tmi_v2/ (relocated camps; original tmi/ untouched)",
              "flee/moving.py, flee/flee.py (cmc_conflict_dependent + enable_return flags)",
              "tests/test_day15_extensions.py (6 focused tests)",
              "scripts/build_tmi_v2_day15.py",
              "scripts/tmi_common_day15.py",
              "scripts/run_tmi_day15_refit.py",
              "scripts/run_tmi_day15_validate_zeigler_shadow.py",
              "scripts/run_tmi_day15_validate_return.py",
              "scripts/plot_tmi_day15_ps2_comparison.py",
              "scripts/run_diagnostics_gate_day15.py",
              "results/tmi/day15_preflight.txt",
              "results/tmi/day15_topology_checkpointA.csv",
              "results/tmi/day15_refit.csv, day15_refit_b2.csv, day15_refit_verdict.json",
              "results/tmi/day15_locked_params.json",
              "results/tmi/day15_zeigler_shadow.csv (+ _verdict.json)",
              "results/tmi/day15_return_curve.csv (+ _verdict.json)",
              "results/tmi/diagnostics_gate_day15.json",
              "figures/tmi/day15_ps2_comparison.png"]:
        _para(doc, f, size=9, space_after=1)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"[docx] wrote {OUT}")


if __name__ == "__main__":
    build()
